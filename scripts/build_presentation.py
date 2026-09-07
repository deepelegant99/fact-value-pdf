import re
from pathlib import Path
from urllib.parse import urlparse
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.util import Inches, Pt

T=Path('Template/Ecology Infographics by Slidesgo.pptx')
M=Path('Content/putnam-fact-value-slides.md')
O=Path('Presentation/putnam-fact-value-slides.pptx')
DARK='202528'; TEXT='3C4246'; MUT='6D7478'; LIGHT='F4F6F5'; MID='DCE2E1'; WHITE='FFFFFF'
TEAL='0B7D80'; BLUE='55B8CF'; GREEN='8BC53F'; YELLOW='F3C746'; ORANGE='F2933A'; PALE=['E7F4F3','EAF5F8','F1F7E8','FFF7DD']
ACC=[TEAL,BLUE,GREEN,ORANGE,YELLOW]; FONT='Arial'
def C(h): return RGBColor.from_string(h)
def md(s):
    s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',s); return s.replace('**','').replace('*','').replace('`','').strip()
def parse():
    out=[]
    for ch in re.split(r'\n---\s*\n',M.read_text(encoding='utf8')):
        m=re.search(r'^##\s+Slide\s+(\d+)\s+[—-]\s+(.+)$',ch,re.M)
        if not m: continue
        d={'n':int(m.group(1)),'title':md(m.group(2)),'bul':[],'par':[],'call':[],'table':[],'urls':[]}
        body=ch[m.end():]
        d['urls']=[u for _,u in re.findall(r'\[([^\]]+)\]\(([^)]+)\)',body)]
        for line in body.splitlines():
            z=line.strip()
            if not z: continue
            if z.startswith('|') and z.endswith('|'):
                cells=[md(x) for x in z.strip('|').split('|')]
                if not all(re.fullmatch(r':?-{3,}:?',x) for x in cells): d['table'].append(cells)
            elif z.startswith('- '): d['bul'].append(md(z[2:]))
            elif z.startswith('**') and z.endswith('**'): d['call'].append(md(z))
            elif z.startswith('*') and z.endswith('*'): pass
            else:
                q=md(z)
                if q.endswith('?') and len(q)<145: d['call'].append(q)
                else: d['par'].append(q)
        out.append(d)
    return sorted(out,key=lambda x:x['n'])
def bg(s):
    f=s.background.fill; f.solid(); f.fore_color.rgb=C(WHITE)
def txt(s,t,x,y,w,h,sz=18,col=TEXT,b=False,al=PP_ALIGN.LEFT,va=MSO_VERTICAL_ANCHOR.TOP):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.margin_left=tf.margin_right=Inches(.04); tf.margin_top=tf.margin_bottom=Inches(.03); tf.vertical_anchor=va
    p=tf.paragraphs[0]; p.alignment=al; r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=C(col); return sh
def box(s,x,y,w,h,fill=LIGHT,line=None):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=C(fill); sh.line.color.rgb=C(line) if line else C(fill); return sh
def header(s,d):
    sec=('ECONOMICS & MEASUREMENT',GREEN) if d['n']<=10 else ('POLITICS, HISTORY & INTERPRETATION',BLUE) if d['n']<=19 else ('PUTNAM’S ARGUMENT',TEAL) if d['n']<=28 else ('STAKES & READING',ORANGE)
    txt(s,sec[0],.55,.2,4,.25,8.5,sec[1],True); txt(s,f"{d['n']:02d}",12,.2,.7,.25,8.5,MUT,True,PP_ALIGN.RIGHT); txt(s,d['title'],.55,.52,12,.62,25 if len(d['title'])<58 else 21,DARK,True,va=MSO_VERTICAL_ANCHOR.MIDDLE)
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.55),Inches(1.15),Inches(1.4),Inches(.06)); sh.fill.solid(); sh.fill.fore_color.rgb=C(sec[1]); sh.line.fill.background()
def foot(s,d):
    hs=[]
    for u in d['urls']:
        h=urlparse(u).netloc.replace('www.','')
        if h and h not in hs: hs.append(h)
    if hs: txt(s,'Sources: '+' · '.join(hs),.55,7.1,10,.16,7,MUT)
def call(s,d,y=6.05,col=TEAL):
    if not d['call']: return
    box(s,.75,y,11.85,.62,col,col); txt(s,d['call'][0],.95,y+.06,11.45,.48,16,WHITE,True,PP_ALIGN.CENTER,MSO_VERTICAL_ANCHOR.MIDDLE)
def card(s,x,y,w,h,title,body,accent=TEAL,i=None,fs=14):
    box(s,x,y,w,h,LIGHT); sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+.18),Inches(y+.14),Inches(.58),Inches(.13)); sh.fill.solid(); sh.fill.fore_color.rgb=C(accent); sh.line.fill.background()
    if h < 1.35:
        if i is not None:
            c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+.18),Inches(y+.34),Inches(.38),Inches(.38)); c.fill.solid(); c.fill.fore_color.rgb=C(accent); c.line.fill.background(); txt(s,str(i),x+.18,y+.35,.38,.32,10,WHITE,True,PP_ALIGN.CENTER,MSO_VERTICAL_ANCHOR.MIDDLE); tx=x+.68
        else: tx=x+.18
        txt(s,title,tx,y+.18,min(2.0,w*.32),.24,9.5,DARK,True,va=MSO_VERTICAL_ANCHOR.MIDDLE)
        txt(s,body,tx,y+.45,w-(tx-x)-.18,h-.52,min(fs,11.5),MUT)
        return
    if i is not None:
        c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+.18),Inches(y+.47),Inches(.46),Inches(.46)); c.fill.solid(); c.fill.fore_color.rgb=C(accent); c.line.fill.background(); txt(s,str(i),x+.18,y+.48,.46,.4,13,WHITE,True,PP_ALIGN.CENTER,MSO_VERTICAL_ANCHOR.MIDDLE); tx=x+.78
    else: tx=x+.18
    txt(s,title,tx,y+.44,w-(tx-x)-.18,.42,12,DARK,True,va=MSO_VERTICAL_ANCHOR.MIDDLE); txt(s,body,x+.18,y+.98,w-.36,max(.25,h-1.1),fs,MUT)
def generic_cards(s,d):
    header(s,d); intro=d['par'][0] if d['par'] else ''
    if intro: txt(s,intro,.82,1.38,11.5,.45,14,MUT,al=PP_ALIGN.CENTER)
    b=d['bul'][:4]
    if len(b)<=3:
        for i,z in enumerate(b): card(s,.8+i*4.1,2.05,3.65,3.25,f'POINT {i+1}',z,ACC[i],i+1,15)
    else:
        for i,z in enumerate(b): card(s,.78+(i%2)*6.08,1.95+(i//2)*1.85,5.65,1.5,f'POINT {i+1}',z,ACC[i],i+1,14)
    if len(d['par'])>1: txt(s,d['par'][-1],1.15,5.55,11,.4,13,MUT,al=PP_ALIGN.CENTER)
    call(s,d)
def compare(s,d):
    header(s,d); items=d['bul'] if len(d['bul'])>=2 else d['par'][:2]
    labels=['DESCRIPTION','EVALUATION'] if d['n']==5 else ['SIDE A','SIDE B']
    for i in range(2): card(s,.85+i*6.05,1.75,5.55,3.65,labels[i],items[i] if i<len(items) else '',[BLUE,ORANGE][i],i+1,17)
    extra=[]
    if len(d['bul'])>2: extra=d['bul'][2:]
    elif len(d['par'])>2: extra=d['par'][2:]
    if extra: txt(s,' · '.join(extra),1.15,5.55,11,.42,13,MUT,al=PP_ALIGN.CENTER)
    call(s,d)
def process(s,d):
    header(s,d); b=d['bul'][:4]; n=len(b)
    for i,z in enumerate(b):
        x=.75+i*(12/n); c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+.65),Inches(2.0),Inches(1.35),Inches(1.35)); c.fill.solid(); c.fill.fore_color.rgb=C(ACC[i]); c.line.fill.background(); txt(s,str(i+1),x+.87,2.31,.9,.55,23,WHITE,True,PP_ALIGN.CENTER,MSO_VERTICAL_ANCHOR.MIDDLE); txt(s,z,x,3.65,12/n-0.1,1.6,14,MUT,al=PP_ALIGN.CENTER)
    if d['par']: txt(s,d['par'][0],1.0,1.42,11.3,.42,14,MUT,al=PP_ALIGN.CENTER)
    call(s,d)
def special(s,d):
    n=d['n']; header(s,d)
    if n==3:
        txt(s,'5%',.8,1.65,3.1,1.25,52,GREEN,True,PP_ALIGN.CENTER,MSO_VERTICAL_ANCHOR.MIDDLE); txt(s,'OUTPUT GROWTH',1.0,2.92,2.7,.35,11,DARK,True,PP_ALIGN.CENTER)
        for i,z in enumerate(d['bul'][1:4]): card(s,4.55,1.58+i*1.25,7.4,1.0,f'CONDITION {i+1}',z,[BLUE,ORANGE,TEAL][i],i+1,13)
    elif n==9:
        txt(s,'99',1.0,1.7,2,1.0,42,TEAL,True); txt(s,'1',10.45,1.7,1.4,1.0,42,ORANGE,True,PP_ALIGN.RIGHT)
        a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1),Inches(2.75),Inches(10.85),Inches(.62)); a.fill.solid(); a.fill.fore_color.rgb=C(TEAL); a.line.fill.background(); b=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(11.85),Inches(2.75),Inches(.11),Inches(.62)); b.fill.solid(); b.fill.fore_color.rgb=C(ORANGE); b.line.fill.background()
        for i,z in enumerate(d['bul'][:3]): card(s,.85+i*4.12,3.75,3.7,1.55,f'ASSUMPTION {i+1}',z,ACC[i],i+1,13)
    elif n==10:
        for j,(lab,a,b) in enumerate([('99 / 1',.99,.01),('50 / 50',.5,.5)]):
            x=.9+j*6.0; card(s,x,1.7,5.45,3.6,lab,'Pareto efficient',[ORANGE,GREEN][j],j+1,17); r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+.55),Inches(3.15),Inches(4.25*a),Inches(.45)); r.fill.solid(); r.fill.fore_color.rgb=C(TEAL); r.line.fill.background(); q=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+.55+4.25*a),Inches(3.15),Inches(4.25*b),Inches(.45)); q.fill.solid(); q.fill.fore_color.rgb=C(ORANGE); q.line.fill.background()
    elif n==14 and d['table']:
        rows=d['table'][1:] if len(d['table'])>1 else d['table']
        for i,r in enumerate(rows[:3]): card(s,.85,1.55+i*1.45,11.65,1.16,['POWER','BELIEF','JUSTIFICATION'][i],r[0],ACC[i],i+1,14); txt(s,r[1] if len(r)>1 else '',8.4,1.89+i*1.45,3.6,.36,12,MUT,True,PP_ALIGN.RIGHT)
    elif n==21:
        for i,z in enumerate(d['bul'][:3]):
            x=1.2+i*3.95; c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(2.0),Inches(2.7),Inches(2.7)); c.fill.solid(); c.fill.fore_color.rgb=C([ORANGE,BLUE,GREEN][i]); c.line.fill.background(); txt(s,z.rstrip('.').upper(),x+.25,2.9,2.2,.55,20,WHITE,True,PP_ALIGN.CENTER,MSO_VERTICAL_ANCHOR.MIDDLE)
        if len(d['par'])>=3: txt(s,d['par'][1]+' '+d['par'][2],2.0,5.1,9.3,.75,14,MUT,al=PP_ALIGN.CENTER)
    elif n==30:
        box(s,.85,1.6,4.15,4.7,DARK,DARK); txt(s,'HILARY\nPUTNAM',1.2,2.0,3.45,.9,25,WHITE,True,PP_ALIGN.CENTER,MSO_VERTICAL_ANCHOR.MIDDLE); txt(s,'The Entanglement\nof Fact and Value',1.15,3.15,3.55,1.0,18,GREEN,True,PP_ALIGN.CENTER,MSO_VERTICAL_ANCHOR.MIDDLE); txt(s,'2002 · Chapter 2 · pp. 28-45',1.15,4.55,3.55,.4,12,'D8E0DE',al=PP_ALIGN.CENTER)
        for i,z in enumerate(d['bul'][:3]): card(s,5.55,2.05+i*1.2,6.65,1.0,f'QUESTION {i+1}',z,[BLUE,GREEN,ORANGE][i],i+1,13)
    else: generic_cards(s,d); return
    if d['par']: txt(s,d['par'][-1],1.15,5.5,11,.4,13,MUT,al=PP_ALIGN.CENTER)
    call(s,d)
def cover(s,d):
    bg(s); txt(s,'FACT / VALUE',.75,.55,3,.35,11,TEAL,True); txt(s,d['title'],.75,1.1,6.1,1.35,35,DARK,True,va=MSO_VERTICAL_ANCHOR.MIDDLE); txt(s,d['par'][0] if d['par'] else '',.78,2.65,5.8,.8,18,MUT)
    for x,y,col,lab in [(7.8,1.55,TEAL,'FACT'),(9.05,2.65,GREEN,'VALUE')]:
        c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(3.3),Inches(3.3)); c.fill.solid(); c.fill.fore_color.rgb=C(PALE[0 if col==TEAL else 2]); c.line.color.rgb=C(col); c.line.width=Pt(2); txt(s,lab,x+.7,y+1.35,1.9,.55,20,col,True,PP_ALIGN.CENTER,MSO_VERTICAL_ANCHOR.MIDDLE)
    txt(s,'ENTANGLEMENT',8.8,3.0,2.3,.35,10,DARK,True,PP_ALIGN.CENTER); txt(s,'Hilary Putnam · The Entanglement of Fact and Value',.78,5.45,5.9,.42,14,DARK,True); call(s,d,6.15,TEAL)
def slide_text(s): return '\n'.join(x.text for x in s.shapes if getattr(x,'has_text_frame',False))
def blank(prs):
    for l in prs.slide_layouts:
        if 'blank' in (l.name or '').lower(): return l
    return prs.slide_layouts[-1]
def credit(prs):
    best=None; score=-1
    for s in prs.slides:
        z=slide_text(s).lower()
        if 'slidesgo' in z:
            sc=sum(k in z for k in ['credits','credit','please keep','attribution','thanks'])
            if sc>score: best,score=s,sc
    return best
def remove(prs,s):
    for q in list(prs.slides._sldIdLst):
        if int(q.id)==int(s.slide_id): prs.part.drop_rel(q.rId); prs.slides._sldIdLst.remove(q); return
def main():
    data=parse(); assert len(data)==30
    prs=Presentation(str(T)); lay=blank(prs); cr=credit(prs)
    for s in list(prs.slides):
        if cr is None or s.slide_id!=cr.slide_id: remove(prs,s)
    families={3:'sp',9:'sp',10:'sp',14:'sp',21:'sp',30:'sp',5:'cmp',6:'cmp',16:'cmp',17:'cmp',20:'cmp',23:'cmp',26:'cmp',7:'proc',8:'proc',19:'proc'}
    for d in data:
        s=prs.slides.add_slide(lay); bg(s)
        if d['n']==1: cover(s,d)
        elif families.get(d['n'])=='sp': special(s,d)
        elif families.get(d['n'])=='cmp': compare(s,d)
        elif families.get(d['n'])=='proc': process(s,d)
        else: generic_cards(s,d)
        foot(s,d)
    if cr:
        for q in list(prs.slides._sldIdLst):
            if int(q.id)==int(cr.slide_id): prs.slides._sldIdLst.remove(q); prs.slides._sldIdLst.append(q); break
    else:
        s=prs.slides.add_slide(lay); bg(s); txt(s,'CREDITS',.8,.8,4,.6,30,DARK,True); txt(s,"Template: 'Ecology Infographics' by Slidesgo",.85,1.8,11,.6,20,TEAL,True); txt(s,'slidesgo.com/theme/ecology-infographics',.85,2.65,11,.5,15,BLUE,True)
    O.parent.mkdir(exist_ok=True); prs.save(str(O)); print(f'Wrote {O} ({len(prs.slides)} slides)')
if __name__=='__main__': main()
